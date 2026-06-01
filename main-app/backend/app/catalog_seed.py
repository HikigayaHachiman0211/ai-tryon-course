from __future__ import annotations

import colorsys
import json
import os
import re
from collections import Counter
from pathlib import Path

import openpyxl
from PIL import Image

BACKEND_ROOT = Path(__file__).resolve().parents[1]
PROJECT_ROOT = BACKEND_ROOT.parent
WORKSPACE_ROOT = PROJECT_ROOT.parent
LEGACY_IMAGE_DIR = PROJECT_ROOT / 'downloaded_jd_images' / '羽绒服_png'

COLOR_KEYWORDS = [
    '岩脊深灰',
    '极晶白',
    '月光白',
    '米白',
    '象牙白',
    '奶白',
    '燕麦',
    '米卡灰',
    '岩层灰',
    '钛灰',
    '钛灰色',
    '浅灰',
    '深灰',
    '雾霾蓝',
    '雾蓝',
    '光影蓝',
    '寂静蓝',
    '墨石蓝',
    '夜泊蓝',
    '夜影蓝',
    '风信蓝',
    '藏青色',
    '藏青',
    '落日橘',
    '骐骥红',
    '白灰雪山',
    '潜水绿色',
    '潜水绿',
    '豆绿色',
    '豆绿',
    '羊绒灰',
    '浅麦芽',
    '浅咖',
    '咖色',
    '卡其',
    '米色',
    '杏色',
    '白色',
    '灰色',
    '黑色',
    '蓝色',
    '绿色',
    '红色',
    '橘色',
    '棕色',
    '青光黑色',
    '基础黑色',
    '极夜黑',
    '曜石黑',
    '幻影黑',
    '幻夜黑',
    '静谧黑',
    '摩卡黑色',
    '摩卡黑',
    '超级黑',
]

STYLE_RULES = [
    ('运动', ['运动', '训练', '跑步', '球员', '安踏', '361°', '阿迪达斯', '耐克', '乔丹', '李宁', '鸿星尔克', '匹克', 'Under Armour']),
    ('户外', ['户外', '登山', '露营', '冲锋', '山系', '雪山', '防水', '防泼水']),
    ('通勤', ['通勤', '简约', '极简', '百搭']),
    ('商务', ['商务', '行政', '爸爸装', '西服领', '毛呢']),
    ('休闲', ['休闲', '日常', '情侣', '校服神器']),
    ('潮流', ['潮流', '时尚', '潮牌', '明星同款']),
    ('工装', ['工装', '巴恩', '山系', '多口袋']),
]

FUNCTION_RULES = [
    ('立领', ['立领']),
    ('连帽', ['连帽']),
    ('可脱卸帽', ['脱卸帽', '可脱卸帽', '可拆卸帽']),
    ('翻领', ['翻领', '西服领']),
    ('圆领', ['圆领']),
    ('防风', ['防风']),
    ('防泼水', ['防泼水']),
    ('防水', ['防水']),
    ('三防', ['三防', '防油', '防污']),
    ('防钻绒', ['防钻绒']),
    ('保暖', ['保暖', '锁温', '御寒', '抗寒']),
    ('轻薄', ['轻薄', '轻量', '轻盈']),
    ('加厚', ['加厚', '厚款', '厚羽']),
    ('抗菌', ['抗菌']),
    ('抗静电', ['抗静电']),
    ('自发热', ['自发热']),
    ('石墨烯', ['石墨烯']),
    ('无缝压胶', ['无缝', '压胶']),
    ('可双面穿', ['双面']),
    ('羽绒内胆', ['内胆']),
    ('三合一', ['三合一']),
    ('两件套', ['两件套']),
    ('高蓬松', ['高蓬', '800蓬', '1000蓬', '高充绒']),
]

BRAND_ALIASES = {
    'MINISO': ['MINISO', '名创优品'],
    '南极人': ['南极人'],
    '雪中飞': ['雪中飞'],
    '波司登': ['波司登', 'BOSIDENG'],
    '阿迪达斯': ['阿迪达斯', 'ADIDAS'],
    '鸭鸭': ['鸭鸭', 'YAYA'],
    '安踏': ['安踏', 'ANTA'],
    '361°': ['361°', '361度'],
    '乔丹': ['乔丹'],
    '李宁': ['李宁', 'LI-NING', 'LINING'],
    '骆驼': ['骆驼', 'CAMEL'],
    '雅鹿': ['雅鹿'],
    '罗蒙': ['罗蒙', 'ROMON'],
    '卡宾': ['卡宾', 'CABBEEN'],
    '安德玛': ['安德玛', 'UNDER ARMOUR', 'UNDERARMOUR'],
    'WASSUP': ['WASSUP'],
    '优衣库': ['优衣库', 'UNIQLO'],
    '鸿星尔克': ['鸿星尔克', 'ERKE'],
    '匹克': ['匹克', 'PEAK'],
    '海澜之家': ['海澜之家', 'HLA'],
    '蕉下': ['蕉下', 'BENEUNDER'],
    'FILA': ['FILA', '斐乐'],
    '杉杉': ['杉杉', 'FIRS'],
    'KAPPA': ['KAPPA', '卡帕'],
    '七匹狼': ['七匹狼', 'SEPTWOLVES'],
    '啄木鸟': ['啄木鸟', 'TUCANO', 'PLOVER'],
    '千仞岗': ['千仞岗', '千仞'],
    '真维斯': ['真维斯', 'JEANSWEST'],
    '迪卡侬': ['迪卡侬', 'DECATHLON'],
    '森马': ['森马', 'SEMIR'],
    '特步': ['特步', 'XTEP'],
    'GXG': ['GXG'],
    '马克华菲': ['马克华菲', 'FAIRWHALE'],
    'PUMA': ['PUMA', '彪马'],
    'NEW BALANCE': ['NEW BALANCE', 'NB'],
    'NIKE': ['NIKE', '耐克'],
    'THE NORTH FACE': ['THE NORTH FACE', '北面', '北脸'],
    '恒源祥': ['恒源祥'],
    '坦博尔': ['坦博尔', 'TANBOER'],
    '花花公子': ['花花公子', 'PLAYBOY'],
    '鄂尔多斯': ['鄂尔多斯'],
    'JEEP': ['JEEP', 'JEEP SPIRIT'],
    '班尼路': ['班尼路', 'BALENO'],
    'KARL LAGERFELD': ['KARL LAGERFELD', '卡尔拉格斐'],
    'NAVIGARE': ['NAVIGARE', '纳维凯尔'],
    '伯希和': ['伯希和', 'PELLIOT'],
    '网易严选': ['网易严选'],
    '君羽': ['君羽', 'JUNYU'],
    '黑冰': ['黑冰', 'BLACKICE', 'BLACK ICE'],
    '回力': ['回力', 'WARRIOR'],
    'SKECHERS': ['SKECHERS', '斯凯奇'],
    'JACK JONES': ['JACK JONES', 'JACKJONES', '杰克琼斯'],
    'I.T': ['I.T'],
    'CLOT': ['CLOT'],
    'INXX': ['INXX'],
    'SAUCONY': ['SAUCONY', '索康尼'],
    'NEPA': ['NEPA'],
    '木林森': ['木林森'],
    '始祖鸟': ['始祖鸟', 'ARCTERYX', "ARC'TERYX"],
    '太平鸟': ['太平鸟', 'PEACEBIRD'],
    'BASIC HOUSE': ['BASIC HOUSE', 'BASICHOUSE', '百家好'],
    '稻草人': ['稻草人', 'MEXICAN'],
    '皮尔卡丹': ['皮尔卡丹', 'PIERRE CARDIN', '皮尔 卡丹'],
    'COS': ['COS'],
    'ZARA': ['ZARA'],
    'HM': ['HM', 'H&M'],
    'MONCLER': ['MONCLER', '盟可睐'],
    '唐狮': ['唐狮', 'TONLION'],
    '美特斯邦威': ['美特斯邦威', 'METERSBONWE'],
    '贵人鸟': ['贵人鸟'],
    '茵曼': ['茵曼', 'INMAN'],
    'SALOMON': ['SALOMON', '萨洛蒙'],
    '凯乐石': ['凯乐石', 'KAILAS'],
    'NBA': ['NBA'],
    'ASICS': ['ASICS', '亚瑟士'],
    '诺诗兰': ['诺诗兰', 'NORTHLAND'],
    'ONITSUKA TIGER': ['ONITSUKA TIGER', '鬼塚虎'],
    'MACKAGE': ['MACKAGE'],
    'HOTSUIT': ['HOTSUIT', '后秀'],
    '三福': ['三福'],
}

BRAND_SUFFIX_PATTERN = re.compile(r'(官方旗舰店|旗舰店|集团旗舰店|官方店|官方奥莱旗舰店|奥莱旗舰店|奥特莱斯|专卖店|专营店|集合店|品牌服饰男装|品牌服饰|品牌专柜|服饰配件|品牌运动|运动户外|精致户外|高端户外|户外放氧计划|男装店|女装店|男装网|男装|女装|中国|店)$')
GENERIC_BRAND_KEYWORDS = ['天猫国际', '天猫超市', '全球探物', '自营', '官方', '旗舰', '店铺', '百亿补贴', '品牌优选', '奥莱', '折扣', '特卖', '专场', '线上商', '外贸仓', '工厂', '直销', '国贸', '官营', '户外', '潮牌', '正品', '代购', '装备', '专柜', '直邮', '体育', '潮流', '保暖', '零售', '一站', '鉴定', '旅游', '小店', '供应商', '商铺', '原创', '运动馆', '清仓', '赔三']

# Store names / non-brand identifiers that should NOT be treated as brands
STORE_NAME_BLACKLIST = {
    'TOPSPORTS', 'YYSPORTS', '胜道', '幸运叶子', '小蔡潮牌', 'BOLM STUDIO',
    'YOUSHU 有术', '有术服饰', 'WARMTREES', 'STAY THE NIGHT', 'SIEBER SHOP',
    '琛琛青少年', '全球超值购', '三彩outlets', '极寒户外羽绒服工厂',
    '坦博尔线上特卖', '回力服饰', '男装服饰官营', '商务大牌轻奢',
    '百亿补贴品牌优选', '品牌优选折扣', '奥莱丨潮牌服饰专场', '奥莱外贸仓',
    '杰雷诺服饰', 'NASA联名潮牌馆', 'NASA潮牌折扣01', 'SssCarry阿叔潮牌',
    '时尚品牌女装潮牌', 'IT潮牌style专购', '屌丝文艺男', '黄某某的小',
    '天猫家享服务', '依依精品女装 折扣', '拓路者户外运动馆', 'skechers运动',
    '恒源祥奥莱国贸', '恒源祥羽绒服饰', 'Semir森马品牌直销',
    'KELME卡尔美夸克', '木林森户外', '棉山下', 'AUSTINBEAUTY',
    '韩国跑腿NY娜娜baby', '喵住', '纤莉秀', 'TRENIJ品牌',
    'JEEPSPIRIT奥莱特卖', 'JEEPSPIRIT弘纳', 'JEEPSPIRIT 品牌',
    'Map Explore 线上商', 'CHENXI HE何晨曦', 'ZHIYAN XU 智研',
    'JEEPSPIRIT时尚', 'JEEPSPIRIT运动服饰', 'JEEPSPIRIT品牌线上',
    'JEEPSPIRIT徽创服饰', 'JEEP SPIRIT正品', 'JEEP SPIRIT 户外服饰',
    'Jeep Spirit 大码男装品牌', 'Jeep SPIRIT大陆总', 'JEEP休闲', 'jeep户外', 'jeep服饰',
    'puma彪马恩树', 'champion运动', 'KELME卡尔美斗格',
    'WILLIAMFOXSONSOUTLET', 'WILLIAM FOX&SONS',
    '回力星驰', '回力休闲服饰', '贵人鸟潮流',
    '拓路者品牌户外', '拓路者世纪弘尚', '拓路者探索者', '拓路者',
    'd[s172667810]', 'K787(原KKYESIOU)', 'MMJ的衣橱', 'E77 GCCDUP',
    '一个强迫症患者', '包oo子和他的朋友们', '你初恋老王ncllw',
    'MADEEXTREME(EME)', '9TH(原VBShowU)', '男道Nandchn',
    '苏宁易购', '银泰百货', '珍品网', '天天洗衣', '立白集团', '丰巢',
    '企鹅', '斯凯奇断码折扣', '李柠羽绒服工厂',
    'Lume Neo 鹅绒羽绒服', '小蚂蚁Ant Studio', 'IFSEN品牌',
    'AsentDboom 轻奢日潮馆', 'SWISS MILITARY功能',
    '迈特优服饰', 'Nova Shadow新星踪迹', '岚室 Clamisgold',
    'rocawear服饰', 'EBLIS HUNGI木乃伊', '鹅绒品牌轻奢 互联网高端',
    '黑青HEIQING STUDIOS', '凯莉欧服饰', '品牌服饰甄选总',
    '奥特莱斯断码特价男装官舰', 'mpf服饰', '田田服饰小屋',
    'VARSDEN 品牌', '名品男装1号', '圆圆品牌精品屋',
}
ACCESSORY_KEYWORDS = ['补丁贴', '修补贴', '布贴', '自粘', '免缝', '修补', '修复', '贴纸', '胶水', '清洗剂', '防尘罩', '拉链', '湿巾', '清洁', '去污', '去渍', '擦鞋', '干洗', '家居', '百货', '代购', 't恤', 'T恤', '背包', '裤子', '罩衣', '鞋', '袜', '帽子', '围巾', '收纳袋', '压缩袋', '收纳神器', '四件套', '被子', '行李箱', '纸扎', '祭祀', '烧纸', '冥币', '寒衣节', '纸衣', '公仔', '整理收纳', '干衣袋', '电吹风机', '吹风机', '除湿', '按扣', '四合扣', '子母扣', '纸棉衣', '纸业', '被套']
GARMENT_KEYWORDS = ['羽绒服', '外套', '大衣', '夹克', '马甲', '背心', '面包服', '内胆', '棉服', '冲锋衣']

def clean_text(raw: object) -> str:
    return re.sub(r'\s+', ' ', str(raw or '').replace('\n', ' ')).strip()

def parse_price(raw: object) -> float | None:
    if raw is None:
        return None
    if isinstance(raw, (int, float)):
        return float(raw)

    text = clean_text(raw)
    if not text:
        return None

    try:
        return float(text)
    except ValueError:
        parts = re.findall(r'\d+(?:\.\d+)?', text)
        if not parts:
            return None
        return float(parts[0])

def color_from_text(text: str) -> str | None:
    for keyword in sorted(set(COLOR_KEYWORDS), key=len, reverse=True):
        if keyword in text:
            return keyword
    match = re.search(
        r'([\u4e00-\u9fa5]{1,8}(?:黑色|白色|灰色|蓝色|绿色|红色|橘色|棕色|咖色|卡其色|米色|米白|杏色|黑|白|灰|蓝|绿|红|橘|棕|咖))',
        text,
    )
    if match:
        return match.group(1)
    return None

def is_skin_like(rgb: tuple[int, int, int]) -> bool:
    r, g, b = rgb
    return r > 135 and g > 95 and b > 70 and r > g > b and (r - b) > 25

def map_rgb_to_color_name(rgb: tuple[int, int, int]) -> str:
    r, g, b = rgb
    h, s, v = colorsys.rgb_to_hsv(r / 255, g / 255, b / 255)

    if v < 0.22:
        return '黑色系'
    if s < 0.10:
        if v > 0.9:
            return '米白色系'
        if v > 0.72:
            return '浅灰色系'
        if v > 0.42:
            return '中灰色系'
        return '深灰色系'

    if 0.08 <= h <= 0.16:
        if v > 0.75:
            return '暖卡其色系'
        return '卡其棕色系'
    if h <= 0.04 or h >= 0.95:
        return '红棕色系' if v < 0.68 else '红色系'
    if 0.20 <= h <= 0.45:
        return '橄榄绿色系'
    if 0.52 <= h <= 0.72:
        return '灰蓝色系' if v > 0.55 else '深蓝色系'
    if r > 160 and g > 140 and b < 120:
        return '暖卡其色系'
    if r > 100 and g > 75 and b < 85:
        return '卡其棕色系'
    if v > 0.78 and s < 0.22:
        return '米白色系'
    return '灰褐色系'

def extract_visual_color(image_path: Path) -> str:
    with Image.open(image_path) as img:
        image = img.convert('RGB').resize((220, 220))

    width, height = image.size
    pixels = image.load()

    corner_points = [
        pixels[0, 0],
        pixels[width - 1, 0],
        pixels[0, height - 1],
        pixels[width - 1, height - 1],
    ]
    bg = tuple(sum(point[channel] for point in corner_points) // 4 for channel in range(3))

    color_scores: Counter[str] = Counter()
    x_start, x_end = int(width * 0.16), int(width * 0.84)
    y_start, y_end = int(height * 0.1), int(height * 0.92)
    for y in range(y_start, y_end):
        for x in range(x_start, x_end):
            r, g, b = pixels[x, y]
            if is_skin_like((r, g, b)):
                continue

            _, s, v = colorsys.rgb_to_hsv(r / 255, g / 255, b / 255)
            distance = ((r - bg[0]) ** 2 + (g - bg[1]) ** 2 + (b - bg[2]) ** 2) ** 0.5
            if distance < 32 and v > 0.45:
                continue
            if v > 0.96 and s < 0.05:
                continue
            color_name = map_rgb_to_color_name((r, g, b))
            center_weight = 1.6 - abs((x / width) - 0.5) - abs((y / height) - 0.48)
            weight = max(center_weight, 0.25)
            if color_name == '黑色系':
                weight += 0.35
            if v < 0.3:
                weight += 0.2
            color_scores[color_name] += weight

    if not color_scores:
        return map_rgb_to_color_name(bg)

    return color_scores.most_common(1)[0][0]

def detect_style_type(text: str) -> str:
    if any(word in text for word in ['马甲', '背心']):
        return '马甲'
    if '派克' in text:
        return '派克大衣'
    if any(word in text for word in ['大衣', '风衣', '毛呢', '西服领', '双面呢']):
        return '中长款大衣'
    if any(word in text for word in ['工装', '巴恩', '山系', '露营', '多口袋']):
        return '巴恩风/工装风'
    if any(word in text for word in ['排骨', '绗缝', '绗线', '小龟背', '内胆']):
        return '绗缝款（排骨款）'
    if any(word in text for word in ['面包', '泡芙', '火山']):
        return '面包服'
    if any(word in text for word in ['过膝', '长款']):
        return '长款'
    if '中长款' in text:
        return '中长款'
    if '短款' in text:
        return '短款'
    if '轻薄' in text:
        return '轻薄款'
    return '常规短外套'

def detect_fit(text: str, style_type: str) -> str:
    if any(word in text for word in ['修身', '收腰']):
        return '修身版，适合纤细到标准身材'
    if any(word in text for word in ['宽松', '廓形', '茧型', '情侣', '潮牌', '假两件']):
        return '宽松廓形，适合标准到微胖身材，利于叠穿'
    if any(word in text for word in ['大码', '爸爸装']):
        return '大码宽松版，适合高壮或微胖身材'
    if style_type == '面包服' or any(word in text for word in ['泡芙', '加厚']):
        return '蓬松宽松版，适合标准到微胖身材，包裹感更强'
    if style_type == '马甲':
        return '标准直筒版，适合叠穿，标准身材友好'
    if style_type in {'长款', '中长款', '中长款大衣', '派克大衣'}:
        return '直筒或微宽松长版，适合高个身材或需要遮臀保暖的人群'
    if any(word in text for word in ['短款', '轻薄']):
        return '标准直筒短版，适合标准身材与小个子日常穿搭'
    return '常规版型，适合大多数标准身材'

def detect_style_features(text: str) -> list[str]:
    features = []
    for label, keywords in STYLE_RULES:
        if any(keyword in text for keyword in keywords):
            features.append(label)
    if not features:
        features.append('休闲')
    return features[:3]

def detect_function_features(text: str) -> list[str]:
    features = []
    for label, keywords in FUNCTION_RULES:
        if any(keyword in text for keyword in keywords):
            features.append(label)
    if not features:
        features.append('基础保暖')
    return features[:6]

def normalize_brand_name(value: str | None) -> str | None:
    text = clean_text(value)
    if not text:
        return None

    # Check blacklist first
    if text in STORE_NAME_BLACKLIST:
        return None

    lowered = text.lower()
    for canonical, aliases in BRAND_ALIASES.items():
        if any(alias.lower() in lowered for alias in aliases):
            return canonical

    cleaned = BRAND_SUFFIX_PATTERN.sub('', text).strip(' -_')
    if not cleaned:
        return None
    if cleaned in STORE_NAME_BLACKLIST:
        return None
    if any(keyword in cleaned for keyword in GENERIC_BRAND_KEYWORDS):
        return None
    if re.fullmatch(r'[A-Za-z0-9& .\'-]+', cleaned):
        return cleaned.upper()
    if len(cleaned) > 18:
        return None
    return cleaned

def infer_brand_name(title: str | None, store_name: str | None = None) -> str | None:
    combined = ' '.join(part for part in [clean_text(title), clean_text(store_name)] if part)
    brand = normalize_brand_name(combined)
    if brand:
        return brand

    title_brand = normalize_brand_name(title)
    if title_brand:
        return title_brand

    return normalize_brand_name(store_name)

def is_valid_taobao_product(title: str, combined_text: str) -> bool:
    if not any(keyword in combined_text for keyword in ['羽绒', '鸭绒', '鹅绒']):
        return False
    if any(keyword in title for keyword in ACCESSORY_KEYWORDS):
        return False
    if not any(keyword in title for keyword in GARMENT_KEYWORDS):
        return False
    return True

def extract_brand_tag(style_features: list[str] | None) -> str | None:
    for feature in style_features or []:
        text = str(feature).strip()
        if text.startswith('品牌:'):
            return text.split(':', 1)[1].strip() or None
        if text.startswith('品牌：'):
            return text.split('：', 1)[1].strip() or None
    return None

def ensure_brand_tag(style_features: list[str] | None, brand: str | None) -> list[str]:
    normalized = [str(feature).strip() for feature in style_features or [] if str(feature).strip()]
    without_brand = [feature for feature in normalized if not feature.startswith(('品牌:', '品牌：'))]
    if brand:
        without_brand.insert(0, f'品牌:{brand}')
    return without_brand

def find_latest_taobao_dataset_dir() -> Path | None:
    env_path = os.getenv('TAOBAO_DATASET_DIR')
    if env_path:
        candidate = Path(env_path)
        if candidate.is_dir():
            return candidate

    candidates = [path for path in WORKSPACE_ROOT.glob('img羽绒服*') if path.is_dir()]
    if not candidates:
        return None

    def sort_key(path: Path) -> tuple[int, float]:
        match = re.search(r'(\d+)$', path.name)
        return (int(match.group(1)) if match else 0, path.stat().st_mtime)

    return sorted(candidates, key=sort_key, reverse=True)[0]

def find_taobao_workbook(dataset_dir: Path | None = None) -> Path | None:
    target_dir = dataset_dir or find_latest_taobao_dataset_dir()
    if target_dir is None:
        return None
    workbooks = sorted(target_dir.glob('*.xlsx'))
    if not workbooks:
        return None
    return workbooks[0]

def build_image_index(dataset_dir: Path) -> dict[str, Path]:
    index: dict[str, Path] = {}
    for image_path in dataset_dir.rglob('*'):
        if image_path.is_file():
            index.setdefault(image_path.name, image_path)
    return index

def resolve_taobao_image_path(row: dict, dataset_dir: Path, image_index: dict[str, Path]) -> Path | None:
    for key in ('商品主图', '图片路径'):
        raw_value = row.get(key)
        if not raw_value:
            continue
        raw_path = Path(str(raw_value))
        if raw_path.is_file():
            return raw_path
        if raw_path.name in image_index:
            return image_index[raw_path.name]
        candidate = dataset_dir / raw_path.name
        if candidate.is_file():
            return candidate
    return None

def load_taobao_seed_rows(workbook_path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [cell for cell in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    dataset_dir = workbook_path.parent
    image_index = build_image_index(dataset_dir)

    rows: list[dict] = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        title = clean_text(row.get('商品名称') or row.get('商品标题'))
        if not title:
            continue

        price = parse_price(row.get('价格'))
        if price is None:
            continue
        if price < 50:
            continue

        detail = clean_text(row.get('商品标题详情'))
        store_name = clean_text(row.get('店铺名称'))
        combined_text = ' '.join(part for part in [title, detail, store_name] if part).strip()
        if not is_valid_taobao_product(title, combined_text):
            continue
        image_path = resolve_taobao_image_path(row, dataset_dir, image_index)
        if image_path is None:
            continue

        style_type = detect_style_type(combined_text)
        brand = infer_brand_name(title, store_name)
        text_color = color_from_text(combined_text)
        rows.append(
            {
                '商品标题': title,
                '价格': price,
                '图片路径': f'{dataset_dir.name}/{image_path.name}',
                '款式类型': style_type,
                '精准颜色色系': text_color or extract_visual_color(image_path),
                '版型与身材适配度': detect_fit(combined_text, style_type),
                '风格特征': ensure_brand_tag(detect_style_features(combined_text), brand),
                '功能属性': detect_function_features(combined_text),
                '品牌': brand,
                '商品链接': clean_text(row.get('商品链接')),
            }
        )

    return rows


def resolve_seed_source(default_database_json: Path, database_json_path: Path | None = None) -> tuple[str, Path]:
    if database_json_path:
        return ('taobao', database_json_path) if database_json_path.suffix.lower() == '.xlsx' else ('json', database_json_path)

    env_path = os.getenv('DATABASE_JSON_PATH')
    if env_path:
        env_seed_path = Path(env_path)
        return ('taobao', env_seed_path) if env_seed_path.suffix.lower() == '.xlsx' else ('json', env_seed_path)

    taobao_workbook = find_taobao_workbook()
    if taobao_workbook is not None:
        return ('taobao', taobao_workbook)

    return ('json', default_database_json)


def load_catalog_seed_rows(default_database_json: Path, database_json_path: Path | None = None) -> list[dict]:
    source_type, data_path = resolve_seed_source(default_database_json, database_json_path)
    if source_type == 'taobao':
        return load_taobao_seed_rows(data_path)
    with data_path.open('r', encoding='utf-8') as file:
        return json.load(file)


def infer_platform_key(image_path: str | None) -> str:
    normalized = str(image_path or '').replace('\\', '/').lstrip('/')
    lowered = normalized.lower()
    if lowered.startswith('downloaded_jd_images/'):
        return 'jd'
    if normalized.startswith('img羽绒服'):
        return 'taobao'
    return 'unknown'


def get_platform_label(platform_key: str | None) -> str:
    labels = {
        'jd': '京东',
        'taobao': '淘宝',
        'unknown': '未标注',
    }
    return labels.get(str(platform_key or '').strip().lower(), '未标注')


def is_within_directory(path: Path, root: Path) -> bool:
    try:
        path.resolve().relative_to(root.resolve())
        return True
    except ValueError:
        return False


def resolve_local_image_path(image_path: str) -> Path | None:
    if not image_path or image_path.startswith(('http://', 'https://')):
        return None

    normalized = image_path.replace('\\', '/').lstrip('/')
    raw_path = Path(normalized)
    taobao_dir = find_latest_taobao_dataset_dir()
    allowed_roots = [root for root in [LEGACY_IMAGE_DIR, taobao_dir] if root and root.exists()]

    candidates: list[Path] = []
    if raw_path.is_absolute():
        candidates.append(raw_path)

    candidates.extend([
        PROJECT_ROOT / normalized,
        WORKSPACE_ROOT / normalized,
        LEGACY_IMAGE_DIR / raw_path.name,
    ])
    if taobao_dir:
        candidates.extend([
            taobao_dir / normalized,
            taobao_dir / raw_path.name,
        ])

    for candidate in candidates:
        try:
            resolved = candidate.resolve()
        except OSError:
            continue
        if not resolved.is_file():
            continue
        if any(is_within_directory(resolved, root) for root in allowed_roots):
            return resolved

    return None
